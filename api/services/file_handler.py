"""
File handler service: Excel upload, sheet selection, normalization, persistence.
"""
import hashlib
import shutil
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import pandas as pd
from datetime import datetime


class FileHandler:
    """Maneja la carga, validación y persistencia de archivos Excel."""
    
    def __init__(self, data_dir: str = "data"):
        self.data_dir = Path(data_dir)
        self.data_dir.mkdir(parents=True, exist_ok=True)
    
    def save_uploaded_file(self, file_bytes: bytes, filename: str) -> Tuple[str, Path]:
        """
        Guarda archivo subido y retorna dataset_id + path.
        
        Args:
            file_bytes: Contenido del archivo
            filename: Nombre original del archivo
            
        Returns:
            (dataset_id, file_path)
        """
        # Generar ID único basado en contenido + timestamp
        content_hash = hashlib.md5(file_bytes).hexdigest()[:16]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dataset_id = f"{content_hash}_{timestamp}"
        
        # Guardar archivo
        file_path = self.data_dir / f"{dataset_id}_{filename}"
        file_path.write_bytes(file_bytes)
        
        return dataset_id, file_path
    
    def get_sheet_names(self, file_path: Path) -> List[str]:
        """
        Lee nombres de hojas del Excel. Para CSV retorna None.
        
        Args:
            file_path: Ruta al archivo Excel o CSV
            
        Returns:
            Lista de nombres de hojas (None si es CSV)
        """
        # Si es CSV, retornar None
        if file_path.suffix.lower() == '.csv':
            return None
        
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            return wb.sheetnames
        except Exception as e:
            # Fallback con pandas
            xl = pd.ExcelFile(file_path)
            return xl.sheet_names
    
    def load_excel(
        self,
        file_path: Path,
        sheet_name: Optional[str] = None
    ) -> Tuple[pd.DataFrame, str]:
        """
        Carga Excel o CSV y normaliza valores faltantes.
        
        Args:
            file_path: Ruta al archivo Excel o CSV
            sheet_name: Nombre de hoja (None = primera hoja o CSV completo)
            
        Returns:
            (DataFrame normalizado, sheet_name utilizado)
        """
        # Si es CSV, leer CSV con dtype especial para columnas ID-like
        if file_path.suffix.lower() == '.csv':
            # Detectar columnas que deben ser string (IDs, celulares, etc.)
            df_sample = pd.read_csv(file_path, nrows=100, encoding='utf-8-sig')
            dtype_dict = {}
            
            for col in df_sample.columns:
                col_lower = col.lower()
                # Forzar string para columnas que son IDs/celulares/documentos
                if any(word in col_lower for word in ['celular', 'telefono', 'phone', 'codigo', 'código', 
                                                       'cedula', 'dni', 'documento', 'identificac', 'tin', 'id']):
                    dtype_dict[col] = str
            
            # Leer con tipos forzados y encoding UTF-8
            df = pd.read_csv(file_path, dtype=dtype_dict, low_memory=False, encoding='utf-8-sig')
            sheet_name = file_path.stem  # Usar nombre del archivo sin extensión
        else:
            # Leer hoja de Excel
            if sheet_name is None:
                sheets = self.get_sheet_names(file_path)
                sheet_name = sheets[0] if sheets else 0
            
            df = pd.read_excel(file_path, sheet_name=sheet_name)
        
        # Normalizar valores faltantes comunes
        df = self._normalize_missing_values(df)
        
        return df, str(sheet_name)
    
    def _normalize_missing_values(self, df: pd.DataFrame) -> pd.DataFrame:
        """
        Normaliza valores que representan missing: '', ' ', 'NA', 'N/A', '-', 'null', etc.
        
        Args:
            df: DataFrame original
            
        Returns:
            DataFrame con valores normalizados a pd.NA
        """
        missing_placeholders = [
            '',
            ' ',
            'NA',
            'N/A',
            'n/a',
            'na',
            '#N/A',
            '#NA',
            '-',
            '--',
            'null',
            'NULL',
            'None',
            'NONE',
            '.',
            '..',
        ]
        
        df_norm = df.copy()
        
        for col in df_norm.columns:
            if df_norm[col].dtype == 'object':  # Solo para strings
                # Reemplazar placeholders
                df_norm[col] = df_norm[col].replace(missing_placeholders, pd.NA)
                
                # Strip whitespace
                try:
                    df_norm[col] = df_norm[col].str.strip()
                    # Reemplazar strings vacíos después del strip
                    df_norm[col] = df_norm[col].replace('', pd.NA)
                except AttributeError:
                    pass
        
        return df_norm
    
    def save_dataframe(self, df: pd.DataFrame, dataset_id: str, format: str = "parquet") -> Path:
        """
        Guarda DataFrame procesado.
        
        Args:
            df: DataFrame a guardar
            dataset_id: ID del dataset
            format: Formato ('parquet', 'csv')
            
        Returns:
            Path al archivo guardado
        """
        if format == "parquet":
            file_path = self.data_dir / f"{dataset_id}.parquet"
            df.to_parquet(file_path, index=False)
        elif format == "csv":
            file_path = self.data_dir / f"{dataset_id}.csv"
            df.to_csv(file_path, index=False)
        else:
            raise ValueError(f"Formato no soportado: {format}")
        
        return file_path
    
    def load_dataframe(self, dataset_id: str) -> Optional[pd.DataFrame]:
        """
        Carga DataFrame guardado (busca parquet primero, luego csv).
        
        Args:
            dataset_id: ID del dataset
            
        Returns:
            DataFrame o None si no existe
        """
        parquet_path = self.data_dir / f"{dataset_id}.parquet"
        if parquet_path.exists():
            return pd.read_parquet(parquet_path)
        
        csv_path = self.data_dir / f"{dataset_id}.csv"
        if csv_path.exists():
            return pd.read_csv(csv_path)
        
        return None
    
    def cleanup(self, dataset_id: str):
        """
        Elimina archivos asociados a un dataset.
        
        Args:
            dataset_id: ID del dataset a limpiar
        """
        for file in self.data_dir.glob(f"{dataset_id}*"):
            file.unlink()
    
    def get_file_info(self, file_path: Path) -> Dict:
        """
        Obtiene información básica del archivo.
        
        Args:
            file_path: Ruta al archivo
            
        Returns:
            Dict con metadata del archivo
        """
        stat = file_path.stat()
        return {
            "filename": file_path.name,
            "size_bytes": stat.st_size,
            "size_mb": round(stat.st_size / (1024 * 1024), 2),
            "created": datetime.fromtimestamp(stat.st_ctime).isoformat(),
            "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(),
        }
